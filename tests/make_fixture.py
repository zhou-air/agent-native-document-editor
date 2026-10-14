from docx import Document
from docx.shared import Pt, Inches, Cm
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw
from pathlib import Path
import json
root=Path(__file__).resolve().parent
(root/'fixtures').mkdir(exist_ok=True)
for name,color,label in [('a','#d9e8f4','P-101'),('b','#e6efdf','V-201')]:
    d=root/'fixtures'/name;d.mkdir(exist_ok=True)
    im=Image.new('RGB',(640,240),'white');draw=ImageDraw.Draw(im)
    draw.rounded_rectangle((35,45,240,190),20,fill=color,outline='#436485',width=4)
    draw.line((240,120,590,120),fill='#436485',width=10)
    draw.text((100,105),label,fill='#24465f',font_size=26)
    draw.text((325,70),'Safety system',fill='#24465f',font_size=26)
    im.save(d/'diagram.png')
d=Document();sec=d.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2);sec.left_margin=Cm(2.8);sec.right_margin=Cm(2.2)
normal=d.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(12);normal._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');normal.paragraph_format.line_spacing=1.5;normal.paragraph_format.first_line_indent=Cm(.74)
for i,size in [(1,20),(2,16),(3,14)]:
    s=d.styles[f'Heading {i}'];s.font.name='Arial';s.font.size=Pt(size);s.font.bold=True;s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(6)
s=d.styles.add_style('EngineeringSubheading',WD_STYLE_TYPE.PARAGRAPH);s.base_style=d.styles['Heading 3']
s=d.styles.add_style('BoldBody',WD_STYLE_TYPE.PARAGRAPH);s.base_style=normal;s.font.bold=True
# Automatic heading numbering, begins at 7, continues at 8.
numbering=d.part.numbering_part.element
abstract=OxmlElement('w:abstractNum');abstract.set(qn('w:abstractNumId'),'77');lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
for tag,value in [('start','7'),('numFmt','decimal'),('lvlText','%1')]:
    el=OxmlElement('w:'+tag);el.set(qn('w:val'),value);lvl.append(el)
abstract.append(lvl);numbering.append(abstract)
num=OxmlElement('w:num');num.set(qn('w:numId'),'77');aid=OxmlElement('w:abstractNumId');aid.set(qn('w:val'),'77');num.append(aid);numbering.append(num)
def numbered(p):
    np=OxmlElement('w:numPr')
    for tag,v in [('ilvl','0'),('numId','77')]:
        x=OxmlElement('w:'+tag);x.set(qn('w:val'),v);np.append(x)
    p._p.get_or_add_pPr().append(np)
d.add_paragraph('1. 普通编号正文，不应成为章节。')
d.add_heading('1 项目概述',1)
p=d.add_paragraph('本项目为测试工程。');p.add_run('设计压力为 1.6 MPa。').bold=True;p.add_run('应设置独立泄压设施。').italic=True
p=d.add_paragraph('继承加粗，',style='BoldBody');p.add_run('此处明确取消粗体。').bold=False
for _ in range(2):d.add_paragraph('相同内容也必须具有不同的块 ID。')
d.add_heading('1.1 设计依据',2)
d.add_paragraph('执行现行标准，保留技术数据与原始编号。')
d.add_paragraph('检查设备位号',style='List Number');d.add_paragraph('检查设计参数',style='List Number');d.add_paragraph('记录核验结果',style='List Bullet')
d.add_paragraph('1.1.1 安全设施清单',style='EngineeringSubheading')
d.add_paragraph('表 1 安全设施与设定值',style='Caption')
t=d.add_table(rows=4,cols=3);t.style='Table Grid'
for row,vals in zip(t.rows,[['设施类别','设备位号','设定值'],['泄压系统','PSV-101','1.6 MPa'],['','PSV-102','1.5 MPa'],['联锁说明','','独立设置']]):
    for cell,value in zip(row.cells,vals):cell.text=value
t.cell(1,0).merge(t.cell(2,0));t.cell(3,0).merge(t.cell(3,1));t.cell(3,0).add_paragraph('禁止省略该说明。')
d.add_paragraph('图 1 安全设施示意',style='Caption');d.add_picture(str(root/'fixtures/a/diagram.png'),width=Inches(5))
d.add_picture(str(root/'fixtures/a/diagram.png'),width=Inches(3))
d.add_picture(str(root/'fixtures/b/diagram.png'),width=Inches(3))
t2=d.add_table(rows=1,cols=1);t2.cell(0,0).paragraphs[0].add_run().add_picture(str(root/'fixtures/b/diagram.png'),width=Inches(2))
p=d.add_heading('应急保障',1);numbered(p);d.add_paragraph('必须保留完整工程说明。')
p=d.add_heading('运行管理',1);numbered(p);d.add_paragraph('关闭后重新打开，数据应完整恢复。')
d.save(root/'fixtures/engineering.docx')
# Hand-authored logical-grid oracle independent of importer.
expected={'titles':['前言','1 项目概述','1.1 设计依据','1.1.1 安全设施清单','7 应急保障','8 运行管理'],'levels':[0,1,2,3,1,1], 'matrix':[['设施类别','设备位号','设定值'],['泄压系统','PSV-101','1.6 MPa'],['泄压系统','PSV-102','1.5 MPa'],['联锁说明\n禁止省略该说明。','联锁说明\n禁止省略该说明。','独立设置']]}
(root/'fixtures/expected.json').write_text(json.dumps(expected,ensure_ascii=False,indent=2))
print(root/'fixtures/engineering.docx')
